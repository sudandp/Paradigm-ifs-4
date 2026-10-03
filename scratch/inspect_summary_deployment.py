import openpyxl

excel_path = r"C:\Users\sudhan\Downloads\Staff Attendance of Serene@BCU - Version_5.7- September 2026 (1).xlsm"
wb = openpyxl.load_workbook(excel_path, data_only=True)

for sname in ["Summary", "Deployment", "Source", "Guide Lines"]:
    if sname not in wb.sheetnames:
        continue
    ws = wb[sname]
    print(f"=== {sname} ===")
    for row_idx, r in enumerate(ws.iter_rows(max_row=30, values_only=True)):
        non_empty = [str(c).strip() for c in r if c is not None and str(c).strip() != '']
        if non_empty:
            print(f"L{row_idx+1}: " + " | ".join(non_empty[:15]))
    print()

wb.close()
