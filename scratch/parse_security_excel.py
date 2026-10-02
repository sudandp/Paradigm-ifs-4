import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\sudhan\Downloads\Version_5.7 SEPTEMBER (1).xlsm', data_only=True)
sheet = wb['Security']

parsed = []
for r in range(7, 7 + 42 * 2, 2):
    bio_id = sheet.cell(r, 3).value
    name = sheet.cell(r, 7).value
    desig = sheet.cell(r, 6).value
    if bio_id is None:
        continue
    
    emp_code = str(bio_id).strip()
    duties = {'double': 0, 'day': 0, 'night': 0, 'wo': 0, 'absent': 0, 'holiday': 0}
    
    for c in range(9, 39):
        d_val = sheet.cell(r, c).value
        n_val = sheet.cell(r+1, c).value
        
        d_str = str(d_val).strip().upper() if d_val else ''
        n_str = str(n_val).strip().upper() if n_val else ''
        
        if d_str == 'P' and n_str == 'P':
            duties['double'] += 1
        elif d_str == 'P':
            duties['day'] += 1
        elif n_str == 'P':
            duties['night'] += 1
        elif d_str == 'W/O' or n_str == 'W/O':
            duties['wo'] += 1
        elif d_str == 'H' or n_str == 'H':
            duties['holiday'] += 1
        else:
            duties['absent'] += 1
            
    parsed.append((emp_code, str(name).strip(), str(desig).strip(), duties))

print(f"Successfully parsed {len(parsed)} security employees:")
for code, name, desig, d in parsed:
    print(f"Code: {code:<10} Name: {name:<22} Double: {d['double']:<2} Day: {d['day']:<2} Night: {d['night']:<2} WO: {d['wo']:<2} Absent: {d['absent']:<2}")
