import openpyxl
import json

excel_path = r"C:\Users\sudhan\Downloads\Staff Attendance of Serene@BCU - Version_5.7- September 2026 (1).xlsm"
wb = openpyxl.load_workbook(excel_path, data_only=True)

dept_sheets = ["Admin", "Electro Mechanical", "HK Services", "Landscaping", "Security"]

inventory = {}

for sname in dept_sheets:
    if sname not in wb.sheetnames:
        continue
    ws = wb[sname]
    inventory[sname] = []
    
    # In each sheet:
    # Header row is usually Row 6:
    # Col A (0): Sl.No
    # Col B (1): Ref No
    # Col C (2): Bio-Metric ID
    # Col D (3): DOJ
    # Col E (4): Staff Category
    # Col F (5): Designation
    # Col G (6): Staff Name
    # Col H (7): Shifts (1, 2, 3, G, or D, N)
    # Col I to AL (8 to 37): Days 1 to 30 of September
    # Col AM onwards: Total Present, WO, Absent, OT, etc.
    
    # Let's inspect headers at row 6
    header_row = [str(c).strip() if c is not None else "" for c in ws[6]]
    
    # Iterate through rows starting row 7
    current_emp = None
    
    for row_idx in range(7, ws.max_row + 1):
        row = [ws.cell(row=row_idx, column=col_idx).value for col_idx in range(1, ws.max_column + 1)]
        
        sl = str(row[0]).strip() if row[0] is not None else ""
        ref = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        bio_id = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
        doj = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ""
        cat = str(row[4]).strip() if len(row) > 4 and row[4] is not None else ""
        desig = str(row[5]).strip() if len(row) > 5 and row[5] is not None else ""
        name = str(row[6]).strip() if len(row) > 6 and row[6] is not None else ""
        shift_code = str(row[7]).strip() if len(row) > 7 and row[7] is not None else ""
        
        # Check if this starts a new employee
        if bio_id and bio_id.lower() not in ['', 'none', 'biometric id', 'bio-metric id']:
            if current_emp:
                inventory[sname].append(current_emp)
            current_emp = {
                'sl': sl,
                'ref': ref,
                'bio_id': bio_id,
                'name': name,
                'desig': desig,
                'category': cat,
                'shifts_data': {}
            }
        
        # If we have a current emp and a shift_code in row[7]
        if current_emp and shift_code:
            # Days 1 to 30 are in columns 9 to 38 (index 8 to 37)
            day_punches = []
            for d in range(1, 31):
                col_i = 7 + d
                val = ws.cell(row=row_idx, column=col_i + 1).value
                val_str = str(val).strip() if val is not None else ""
                day_punches.append(val_str)
            current_emp['shifts_data'][shift_code] = day_punches

    if current_emp:
        inventory[sname].append(current_emp)

print("=== INVENTORY SUMMARY ===")
total_all = 0
for sname, emps in inventory.items():
    print(f"Department: {sname:20} -> Total Employees: {len(emps)}")
    total_all += len(emps)
    for e in emps[:3]:
        print(f"   BioId: {e['bio_id']:8} Name: {e['name']:25} Desig: {e['desig']:25} Shifts: {list(e['shifts_data'].keys())}")

print(f"\nTOTAL ALL EMPLOYEES ACROSS ALL DEPARTMENTS: {total_all}")

# Save detailed JSON for analysis
with open(r"e:\backup\onboarding all files\Paradigm Office 4\scratch\serene_all_staff.json", "w") as f:
    json.dump(inventory, f, indent=2)

print("Saved detailed inventory to scratch/serene_all_staff.json")
wb.close()
