import openpyxl
import os

master_path = r"C:\Users\sudhan\Downloads\Staff Attendance of Serene@BCU - Version_5.7- September 2026 (1).xlsm"
issues_path = r"C:\Users\sudhan\Downloads\Attendane Issues.xlsx"

wb_master = openpyxl.load_workbook(master_path, data_only=True)
wb_issues = openpyxl.load_workbook(issues_path, data_only=True)

# 1. Read HK staff weekly off days from Master Excel HK Services
ws_hk = wb_master['HK Services']
hk_staff = []
for r in range(7, ws_hk.max_row + 1):
    c_bio = ws_hk.cell(r, 3).value
    c_name = ws_hk.cell(r, 7).value
    c_desig = ws_hk.cell(r, 6).value
    c_shift = ws_hk.cell(r, 8).value
    if c_bio and str(c_bio).strip() not in ['', 'None', 'Bio-Metric ID']:
        bio_id = str(c_bio).strip()
        if 'PARVATHI' in str(c_name or '').upper():
            bio_id = '31109-HK'
        wo_days = []
        for d in range(1, 31):
            v = ws_hk.cell(r, 8 + d).value
            if v and str(v).strip().upper() in ['W/O', 'WO']:
                wo_days.append(d)
        hk_staff.append({
            'bio_id': bio_id,
            'name': str(c_name or '').strip(),
            'desig': str(c_desig or '').strip(),
            'wo_days': wo_days
        })

print(f"Loaded {len(hk_staff)} Housekeeping staff.")
for h in hk_staff[:6]:
    print(f"  {h['bio_id']} ({h['name']}): WO days = {h['wo_days']}")
