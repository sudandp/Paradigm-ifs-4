import openpyxl
from collections import Counter

excel_path = r"C:\Users\sudhan\Downloads\Staff Attendance of Serene@BCU - Version_5.7- September 2026 (1).xlsm"
wb = openpyxl.load_workbook(excel_path, data_only=True)

dept_configs = [
    {"sheet": "Admin", "dept_name": "Administration"},
    {"sheet": "Electro Mechanical", "dept_name": "Electro Mechanical"},
    {"sheet": "HK Services", "dept_name": "Housekeeping"},
    {"sheet": "Landscaping", "dept_name": "Landscaping"},
    {"sheet": "Security", "dept_name": "Security"}
]

all_records = []

for cfg in dept_configs:
    sname = cfg["sheet"]
    if sname not in wb.sheetnames:
        continue
    ws = wb[sname]
    
    current_emp = None
    
    for row_idx in range(7, ws.max_row + 1):
        c_bio = ws.cell(row=row_idx, column=3).value
        bio_id = str(c_bio).strip() if c_bio is not None else ""
        c_name = ws.cell(row=row_idx, column=7).value
        name = str(c_name).strip() if c_name is not None else ""
        c_desig = ws.cell(row=row_idx, column=6).value
        desig = str(c_desig).strip() if c_desig is not None else ""
        c_shift = ws.cell(row=row_idx, column=8).value
        shift_tag = str(c_shift).strip() if c_shift is not None else ""
        
        # Security overrides
        if cfg["dept_name"] == "Security":
            if "MRITYUNJAY" in name.upper():
                bio_id = "32085"
            elif "RAM CHANDRA" in name.upper() and "SUPERVISOR" in desig.upper():
                continue
        elif cfg["dept_name"] == "Housekeeping":
            if "PARVATHI" in name.upper():
                bio_id = "31109-HK"
        
        if bio_id and bio_id.lower() not in ['', 'none', 'bio-metric id', 'biometric id', 'bio id']:
            if current_emp:
                all_records.append(current_emp)
            current_emp = {
                "dept": cfg["dept_name"],
                "bio_id": bio_id,
                "name": name,
                "desig": desig,
                "shifts": {}
            }
        
        if current_emp and shift_tag:
            days = []
            for d in range(1, 31):
                col_i = 8 + d
                val = ws.cell(row=row_idx, column=col_i).value
                val_str = str(val).strip() if val is not None else ""
                days.append(val_str)
            current_emp["shifts"][shift_tag] = days

    if current_emp:
        all_records.append(current_emp)

wb.close()

# Calculate mandays for each employee
total_present = 0.0
total_wo = 0
total_absent = 0.0
total_ot_mins = 0

for emp in all_records:
    # merge days
    for d in range(30):
        # find what they worked
        day_tag = ""
        p_val = 0.0
        wo_val = 0
        a_val = 0.0
        ot_mins = 0
        
        if emp["dept"] == "Security":
            d_tag = emp["shifts"].get("D", [""]*30)[d] if "D" in emp["shifts"] else ""
            n_tag = emp["shifts"].get("N", [""]*30)[d] if "N" in emp["shifts"] else ""
            d_p = d_tag.upper() == 'P'
            n_p = n_tag.upper() == 'P'
            if d_p and n_p:
                p_val = 2.0
                ot_mins = 720
            elif d_p or n_p:
                p_val = 1.0
            elif d_tag.upper() == 'WO' or n_tag.upper() == 'WO':
                wo_val = 1
            elif d_tag.upper() == 'A' or n_tag.upper() == 'A':
                a_val = 1.0
            else:
                a_val = 1.0
        else:
            # For other depts, take first shift tag that has an entry
            found = False
            for s_tag, days in emp["shifts"].items():
                tag = days[d].upper() if d < len(days) else ""
                if tag in ['P', 'WO', 'A', 'H', 'PH']:
                    found = True
                    if tag == 'P':
                        p_val = 1.0
                    elif tag == 'WO':
                        wo_val = 1
                    elif tag in ['H', 'PH']:
                        p_val = 1.0
                    else:
                        a_val = 1.0
                    break
            if not found:
                a_val = 1.0
                
        total_present += p_val
        total_wo += wo_val
        total_absent += a_val
        total_ot_mins += ot_mins

print(f"Total Staff: {len(all_records)}")
print(f"Total Present Mandays: {total_present}")
print(f"Total Weekly Offs: {total_wo}")
print(f"Total Absent Days: {total_absent}")
print(f"Total OT Hours: {total_ot_mins / 60.0}")
