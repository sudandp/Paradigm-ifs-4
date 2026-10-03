import openpyxl
import json

excel_path = r"C:\Users\sudhan\Downloads\Staff Attendance of Serene@BCU - Version_5.7- September 2026 (1).xlsm"
wb = openpyxl.load_workbook(excel_path, data_only=True)

dept_configs = [
    {"sheet": "Admin", "dept_name": "Administration", "default_shift": "GS"},
    {"sheet": "Electro Mechanical", "dept_name": "Electro Mechanical", "default_shift": "Rotational"},
    {"sheet": "HK Services", "dept_name": "Housekeeping", "default_shift": "HK-M"},
    {"sheet": "Landscaping", "dept_name": "Landscaping", "default_shift": "GAR"},
    {"sheet": "Security", "dept_name": "Security", "default_shift": "Security-12h"}
]

all_data = []

for cfg in dept_configs:
    sname = cfg["sheet"]
    if sname not in wb.sheetnames:
        continue
    ws = wb[sname]
    
    current_emp = None
    
    for row_idx in range(7, ws.max_row + 1):
        # Bio-Metric ID is col C (3)
        c_bio = ws.cell(row=row_idx, column=3).value
        bio_id = str(c_bio).strip() if c_bio is not None else ""
        
        c_name = ws.cell(row=row_idx, column=7).value
        name = str(c_name).strip() if c_name is not None else ""
        
        c_desig = ws.cell(row=row_idx, column=6).value
        desig = str(c_desig).strip() if c_desig is not None else ""
        
        c_shift = ws.cell(row=row_idx, column=8).value
        shift_tag = str(c_shift).strip() if c_shift is not None else ""
        
        # Check if new employee
        if bio_id and bio_id.lower() not in ['', 'none', 'bio-metric id', 'biometric id', 'bio id']:
            if current_emp:
                all_data.append(current_emp)
            current_emp = {
                "dept": cfg["dept_name"],
                "sheet": sname,
                "bio_id": bio_id,
                "name": name,
                "desig": desig,
                "shifts": {}
            }
        
        if current_emp and shift_tag:
            days = []
            for d in range(1, 31):
                col_i = 8 + d
                val = ws.cell(row=row_idx, column=col_i).value
                val_str = str(val).strip() if val is not None else ""
                days.append(val_str)
            current_emp["shifts"][shift_tag] = days

    if current_emp:
        all_data.append(current_emp)

wb.close()

print(f"Total Parsed Staff: {len(all_data)}")

# Calculate department stats
dept_counts = {}
for e in all_data:
    d = e["dept"]
    dept_counts[d] = dept_counts.get(d, 0) + 1

for d, count in dept_counts.items():
    print(f"  {d:20}: {count} employees")

# Sample check on Vedamurthy and Chandrappa
for e in all_data:
    if e["bio_id"] in ["31001", "31014", "31013", "31016", "32001"]:
        print(f"\nStaff: {e['bio_id']} - {e['name']} ({e['desig']} in {e['dept']})")
        for s_tag, d_vals in e["shifts"].items():
            non_empty_days = [(idx+1, v) for idx, v in enumerate(d_vals) if v]
            print(f"   Shift '{s_tag}': {len(non_empty_days)} active days -> {non_empty_days[:8]}...")
