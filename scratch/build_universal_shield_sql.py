import openpyxl
import os

master_path = r"C:\Users\sudhan\Downloads\Staff Attendance of Serene@BCU - Version_5.7- September 2026 (1).xlsm"
issues_path = r"C:\Users\sudhan\Downloads\Attendane Issues.xlsx"

wb_master = openpyxl.load_workbook(master_path, data_only=True)
wb_issues = openpyxl.load_workbook(issues_path, data_only=True)

# 1. Parse HK Services to get exact shifts for all 26 cleaners
ws_hk = wb_master['HK Services']
hk_records = []
for r in range(7, ws_hk.max_row + 1):
    c_bio = ws_hk.cell(r, 3).value
    c_name = ws_hk.cell(r, 7).value
    c_desig = ws_hk.cell(r, 6).value
    c_shift = ws_hk.cell(r, 8).value
    if c_bio and str(c_bio).strip() not in ['', 'None', 'Bio-Metric ID']:
        bio_id = str(c_bio).strip()
        if 'PARVATHI' in str(c_name or '').upper():
            bio_id = '31109-HK'
        days_shifts = []
        for d in range(1, 31):
            val = ws_hk.cell(r, 8 + d).value
            days_shifts.append(str(val).strip().upper() if val is not None else "")
        hk_records.append({
            'dept': 'Housekeeping',
            'bio_id': bio_id,
            'name': str(c_name or '').strip(),
            'desig': str(c_desig or '').strip(),
            'days': days_shifts
        })

print(f"Loaded {len(hk_records)} Housekeeping staff.")

# 2. Parse Electro Mechanical
ws_em = wb_master['Electro Mechanical']
em_records = []
for r in range(7, ws_em.max_row + 1):
    c_bio = ws_em.cell(r, 3).value
    c_name = ws_em.cell(r, 7).value
    c_desig = ws_em.cell(r, 6).value
    c_shift = ws_em.cell(r, 8).value
    if c_bio and str(c_bio).strip() not in ['', 'None', 'Bio-Metric ID']:
        bio_id = str(c_bio).strip()
        # Find which shift lines this employee has
        # Each employee in EM has multiple shift rows: G, 1, 2, 3
        # Let's inspect how EM is structured
        pass

wb_master.close()
wb_issues.close()
