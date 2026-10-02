import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\sudhan\Downloads\Version_5.7 SEPTEMBER (1).xlsm', data_only=True)
sheet = wb['Security']

target_guards = {
    '320883': ('32088', 9358, 'Kausar Alam'),
    '22082': ('32082', 9226, 'Khakendra'),
    '22079': ('32079', 9228, 'Tularaj kaila'),
    '32086_MRIT': ('32085', 9273, 'Mritunjykumar')
}

patch_records = []

for r in range(7, 7 + 42 * 2, 2):
    bio_id = sheet.cell(r, 3).value
    name = sheet.cell(r, 7).value
    if bio_id is None:
        continue
    
    emp_code = str(bio_id).strip()
    clean_name = str(name).strip()
    
    key = None
    if emp_code == '320883':
        key = '320883'
    elif emp_code == '22082':
        key = '22082'
    elif emp_code == '22079':
        key = '22079'
    elif emp_code == '32086' and 'MRITYUNJAY' in clean_name.upper():
        key = '32086_MRIT'
        
    if not key:
        continue
        
    db_code, db_emp_id, db_name = target_guards[key]
    
    for c in range(9, 39):
        day_num = c - 8
        date_str = f'2026-09-{day_num:02d}'
        d_val = sheet.cell(r, c).value
        n_val = sheet.cell(r+1, c).value
        
        d_str = str(d_val).strip().upper() if d_val else ''
        n_str = str(n_val).strip().upper() if n_val else ''
        
        if d_str == 'P' and n_str == 'P':
            duty_type = 'DOUBLE'
        elif d_str == 'P':
            duty_type = 'DAY'
        elif n_str == 'P':
            duty_type = 'NIGHT'
        elif d_str == 'W/O' or n_str == 'W/O':
            duty_type = 'WO'
        elif d_str == 'H' or n_str == 'H':
            duty_type = 'H'
        else:
            duty_type = 'A'
            
        patch_records.append((db_code, db_emp_id, db_name, date_str, duty_type))

print(f'Extracted {len(patch_records)} patch records for 4 guards.')

sql_header = """-- ====================================================================================================
-- QUICK PATCH FOR REMAINING 4 SECURITY GUARDS (1-SECOND EXECUTION)
-- Database: etimetracklite1
-- Verified DB IDs:
--   1. 32088 (Kausar Alam - EmployeeId 9358)
--   2. 32082 (Khakendra - EmployeeId 9226)
--   3. 32079 (Tularaj kaila - EmployeeId 9228)
--   4. 32085 (Mritunjykumar - EmployeeId 9273)
-- ====================================================================================================

USE [etimetracklite1];
GO
SET NOCOUNT ON;

PRINT '>>> Applying Instant Patch for the 4 Security Guards...';

IF OBJECT_ID('tempdb..#GuardPatch', 'U') IS NOT NULL DROP TABLE #GuardPatch;
CREATE TABLE #GuardPatch (
    EmployeeCode VARCHAR(50),
    EmployeeId INT,
    StaffName VARCHAR(100),
    DutyDate DATE,
    DutyType VARCHAR(20)
);

INSERT INTO #GuardPatch (EmployeeCode, EmployeeId, StaffName, DutyDate, DutyType) VALUES
"""

val_rows = []
for p in patch_records:
    val_rows.append(f"('{p[0]}', {p[1]}, '{p[2]}', '{p[3]}', '{p[4]}')")

sql_body = ",\n".join(val_rows) + ";\n\n"

sql_footer = """-- 1. Insert into dbo.ExcelMuster_Security_Sep2026
INSERT INTO dbo.ExcelMuster_Security_Sep2026 (EmployeeCode, StaffName, Designation, DutyDate, DutyType, DayVal, NightVal)
SELECT p.EmployeeCode, p.StaffName, 'SECURITY GUARD', p.DutyDate, p.DutyType, '', ''
FROM #GuardPatch p
WHERE NOT EXISTS (
    SELECT 1 FROM dbo.ExcelMuster_Security_Sep2026 m 
    WHERE m.EmployeeCode = p.EmployeeCode AND m.DutyDate = p.DutyDate
);

-- 2. Update dbo.AttendanceLogs
UPDATE a
SET 
    a.Duration = CASE WHEN p.DutyType IN ('DAY', 'NIGHT', 'DOUBLE') THEN 720 ELSE 0 END,
    a.OverTime = CASE WHEN p.DutyType = 'DOUBLE' THEN 720 ELSE 0 END,
    a.LateBy = 0,
    a.EarlyBy = 0,
    a.ShiftId = CASE 
        WHEN p.DutyType IN ('DAY', 'DOUBLE') THEN 44
        WHEN p.DutyType = 'NIGHT' THEN 45
        WHEN p.DutyType = 'WO' THEN 1
        WHEN p.DutyType = 'H'  THEN 4
        ELSE 3
    END,
    a.Status = CASE 
        WHEN p.DutyType IN ('DAY', 'NIGHT', 'DOUBLE') THEN 'Present '
        WHEN p.DutyType = 'WO' THEN 'WeeklyOff'
        WHEN p.DutyType = 'H'  THEN 'Holiday'
        ELSE 'Absent'
    END,
    a.StatusCode = CASE 
        WHEN p.DutyType IN ('DAY', 'NIGHT', 'DOUBLE') THEN 'P'
        WHEN p.DutyType = 'WO' THEN 'WO'
        WHEN p.DutyType = 'H'  THEN 'H'
        ELSE 'A'
    END,
    a.P1Status = CASE 
        WHEN p.DutyType IN ('DAY', 'NIGHT', 'DOUBLE') THEN 'P'
        WHEN p.DutyType = 'WO' THEN 'WO'
        WHEN p.DutyType = 'H'  THEN 'H'
        ELSE 'A'
    END,
    a.Present = CASE WHEN p.DutyType IN ('DAY', 'NIGHT', 'DOUBLE') THEN 1.0 ELSE 0.0 END,
    a.Absent  = CASE WHEN p.DutyType = 'A' THEN 1.0 ELSE 0.0 END,
    a.WeeklyOff = CASE WHEN p.DutyType = 'WO' THEN 1 ELSE 0 END,
    a.Holiday   = CASE WHEN p.DutyType = 'H'  THEN 1 ELSE 0 END,
    a.OutTime = CASE 
        WHEN p.DutyType = 'DAY' THEN DATEADD(HOUR, 20, CAST(CAST(a.AttendanceDate AS DATE) AS DATETIME))
        WHEN p.DutyType IN ('NIGHT', 'DOUBLE') THEN DATEADD(HOUR, 32, CAST(CAST(a.AttendanceDate AS DATE) AS DATETIME))
        ELSE '1900-01-01 00:00:00'
    END,
    a.InTime = CASE 
        WHEN p.DutyType IN ('A', 'WO', 'H') THEN '1900-01-01 00:00:00'
        WHEN p.DutyType IN ('DAY', 'DOUBLE') THEN 
            CASE WHEN a.InTime = '1900-01-01 00:00:00' OR a.InTime IS NULL 
                 THEN DATEADD(HOUR, 8, CAST(CAST(a.AttendanceDate AS DATE) AS DATETIME)) 
                 ELSE a.InTime END
        WHEN p.DutyType = 'NIGHT' THEN 
            CASE WHEN a.InTime = '1900-01-01 00:00:00' OR a.InTime IS NULL 
                 THEN DATEADD(HOUR, 20, CAST(CAST(a.AttendanceDate AS DATE) AS DATETIME)) 
                 ELSE a.InTime END
        ELSE a.InTime
    END,
    a.Remarks = 'Excel Verified: ' + p.DutyType
FROM dbo.AttendanceLogs a
JOIN #GuardPatch p 
    ON a.EmployeeId = p.EmployeeId 
   AND CAST(a.AttendanceDate AS DATE) = p.DutyDate;

-- 3. Insert any missing days into dbo.AttendanceLogs
INSERT INTO dbo.AttendanceLogs (
    AttendanceDate, EmployeeId, InTime, OutTime, Duration, LateBy, EarlyBy, 
    IsOnLeave, WeeklyOff, Holiday, PunchRecords, ShiftId, Present, Absent, 
    [Status], StatusCode, P1Status, IsonSpecialOff, OverTime, OverTimeE, 
    MissedOutPunch, MissedInPunch, Remarks
)
SELECT 
    CAST(p.DutyDate AS DATETIME),
    p.EmployeeId,
    CASE 
        WHEN p.DutyType IN ('DAY', 'DOUBLE') THEN DATEADD(HOUR, 8, CAST(p.DutyDate AS DATETIME))
        WHEN p.DutyType = 'NIGHT'  THEN DATEADD(HOUR, 20, CAST(p.DutyDate AS DATETIME))
        ELSE '1900-01-01 00:00:00'
    END,
    CASE 
        WHEN p.DutyType = 'DAY'    THEN DATEADD(HOUR, 20, CAST(p.DutyDate AS DATETIME))
        WHEN p.DutyType IN ('NIGHT', 'DOUBLE') THEN DATEADD(HOUR, 32, CAST(p.DutyDate AS DATETIME))
        ELSE '1900-01-01 00:00:00'
    END,
    CASE WHEN p.DutyType IN ('DAY', 'NIGHT', 'DOUBLE') THEN 720 ELSE 0 END,
    0, 0, 0,
    CASE WHEN p.DutyType = 'WO' THEN 1 ELSE 0 END,
    CASE WHEN p.DutyType = 'H'  THEN 1 ELSE 0 END,
    '',
    CASE 
        WHEN p.DutyType IN ('DAY', 'DOUBLE') THEN 44
        WHEN p.DutyType = 'NIGHT' THEN 45
        WHEN p.DutyType = 'WO' THEN 1
        WHEN p.DutyType = 'H'  THEN 4
        ELSE 3
    END,
    CASE WHEN p.DutyType IN ('DAY', 'NIGHT', 'DOUBLE') THEN 1.0 ELSE 0.0 END,
    CASE WHEN p.DutyType = 'A' THEN 1.0 ELSE 0.0 END,
    CASE 
        WHEN p.DutyType IN ('DAY', 'NIGHT', 'DOUBLE') THEN 'Present '
        WHEN p.DutyType = 'WO' THEN 'WeeklyOff'
        WHEN p.DutyType = 'H'  THEN 'Holiday'
        ELSE 'Absent'
    END,
    CASE 
        WHEN p.DutyType IN ('DAY', 'NIGHT', 'DOUBLE') THEN 'P'
        WHEN p.DutyType = 'WO' THEN 'WO'
        WHEN p.DutyType = 'H'  THEN 'H'
        ELSE 'A'
    END,
    CASE 
        WHEN p.DutyType IN ('DAY', 'NIGHT', 'DOUBLE') THEN 'P'
        WHEN p.DutyType = 'WO' THEN 'WO'
        WHEN p.DutyType = 'H'  THEN 'H'
        ELSE 'A'
    END,
    0,
    CASE WHEN p.DutyType = 'DOUBLE' THEN 720 ELSE 0 END,
    0, 0, 0,
    'Excel Inserted: ' + p.DutyType
FROM #GuardPatch p
WHERE NOT EXISTS (
    SELECT 1 FROM dbo.AttendanceLogs a
    WHERE a.EmployeeId = p.EmployeeId AND CAST(a.AttendanceDate AS DATE) = p.DutyDate
);

-- 4. Sync dbo.EmployeeShiftSchedule roster
DELETE ess
FROM dbo.EmployeeShiftSchedule ess
JOIN #GuardPatch p ON ess.EmployeeId = p.EmployeeId AND CAST(ess.ShiftDate AS DATE) = p.DutyDate;

INSERT INTO dbo.EmployeeShiftSchedule (EmployeeId, ShiftDate, ShiftId)
SELECT 
    p.EmployeeId,
    CAST(p.DutyDate AS DATETIME),
    CASE 
        WHEN p.DutyType IN ('DAY', 'DOUBLE') THEN 44
        WHEN p.DutyType = 'NIGHT' THEN 45
        WHEN p.DutyType = 'WO' THEN 1
        WHEN p.DutyType = 'H'  THEN 4
        ELSE 3
    END
FROM #GuardPatch p;

PRINT '✓ All 4 Guards Updated Successfully in dbo.AttendanceLogs & dbo.EmployeeShiftSchedule!';

-- 5. Verification of the 4 Guards
SELECT 
    e.EmployeeCode,
    e.EmployeeName,
    COUNT(CASE WHEN p.DutyType = 'DOUBLE' THEN 1 END) AS Excel_Double,
    COUNT(CASE WHEN p.DutyType = 'DAY' THEN 1 END) AS Excel_Day,
    COUNT(CASE WHEN p.DutyType = 'NIGHT' THEN 1 END) AS Excel_Night,
    COUNT(CASE WHEN p.DutyType = 'WO' THEN 1 END) AS Excel_WO,
    COUNT(CASE WHEN p.DutyType = 'A' THEN 1 END) AS Excel_Absent,
    CAST(ROUND(ISNULL(SUM(a.OverTime), 0) / 60.0, 1) AS NUMERIC(10,1)) AS Total_OT_Hours_DB,
    CAST(ISNULL(SUM(a.Present), 0) AS NUMERIC(10,1)) AS Total_Present_Days_DB
FROM #GuardPatch p
JOIN dbo.Employees e ON p.EmployeeId = e.EmployeeId
LEFT JOIN dbo.AttendanceLogs a 
    ON a.EmployeeId = p.EmployeeId AND CAST(a.AttendanceDate AS DATE) = p.DutyDate
GROUP BY e.EmployeeCode, e.EmployeeName
ORDER BY e.EmployeeName;
GO
"""

with open(r'e:\backup\onboarding all files\Paradigm Office 4\sql\patch_remaining_4_security_guards.sql', 'w', encoding='utf-8') as f:
    f.write(sql_header + sql_body + sql_footer)

print('Generated sql/patch_remaining_4_security_guards.sql successfully!')
