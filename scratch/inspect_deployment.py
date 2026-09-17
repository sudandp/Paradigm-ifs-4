import openpyxl, difflib, json, urllib.request
from dotenv import dotenv_values

config = dotenv_values('.env.local')
url = config.get('VITE_SUPABASE_URL')
key = config.get('SUPABASE_SERVICE_ROLE_KEY') or config.get('VITE_SUPABASE_SERVICE_ROLE_KEY') or config.get('VITE_SUPABASE_ANON_KEY')

req = urllib.request.Request(f'{url}/rest/v1/organizations?select=id,short_name,full_name,manpower_approved_count', headers={
    'apikey': key,
    'Authorization': f'Bearer {key}'
})

with urllib.request.urlopen(req) as resp:
    db_orgs = json.loads(resp.read().decode('utf-8'))

wb = openpyxl.load_workbook(r'C:\Users\sudhan\Downloads\Version_5.6 Final.xlsm', data_only=True)
sheet = wb['Deployment']

excel_sites = []
for c in range(2, sheet.max_column + 1):
    name = sheet.cell(2, c).value
    tot = sheet.cell(1, c).value
    if name:
        excel_sites.append({'name': str(name).strip(), 'total': float(tot or 0), 'col': c})

print(f'Total DB orgs: {len(db_orgs)}')
print(f'Total Excel sites: {len(excel_sites)}')

# Check matching
matched = []
unmatched = []

for org in db_orgs:
    s_name = org['short_name'] or ''
    f_name = org.get('full_name') or ''
    found = None
    # 1. exact match
    for e in excel_sites:
        if s_name.upper().strip() == e['name'].upper().strip() or f_name.upper().strip() == e['name'].upper().strip():
            found = e
            break
    # 2. fuzzy match
    if not found:
        matches = difflib.get_close_matches(s_name.upper().strip(), [e['name'].upper().strip() for e in excel_sites], n=1, cutoff=0.6)
        if matches:
            for e in excel_sites:
                if e['name'].upper().strip() == matches[0]:
                    found = e
                    break

    if found:
        matched.append((org, found))
    else:
        unmatched.append(org)

print(f'Matched DB orgs: {len(matched)}')
print(f'Unmatched DB orgs: {len(unmatched)}')

print('\nSample matched:')
for org, e in matched[:30]:
    print(f"{org['short_name']} -> {e['name']} : {e['total']}")

print('\nSample unmatched DB orgs:')
for org in unmatched[:30]:
    print(f"{org['short_name']} ({org['id']})")
